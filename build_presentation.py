"""Builds the management presentation from the notebook outputs.

Run the notebook first (it writes figures/ and results/key_numbers.json), then:
    python build_presentation.py
"""
import json
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

K = json.load(open('results/key_numbers.json', encoding='utf-8'))
OUT = 'Customer_Review_Text_Mining.pptx'

NAVY = RGBColor(0x1F, 0x2A, 0x44)
INK = RGBColor(0x1B, 0x1B, 0x1B)
INK2 = RGBColor(0x55, 0x57, 0x5E)
MUTED = RGBColor(0x8A, 0x8C, 0x92)
NEG = RGBColor(0xEB, 0x68, 0x34)
POS = RGBColor(0x2A, 0x78, 0xD6)
TINT = RGBColor(0xF1, 0xF3, 0xF7)
NEG_TINT = RGBColor(0xFD, 0xEE, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = 'Calibri'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5


def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, spacing=None):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        parts = para if isinstance(para, list) else [(para, {})]
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.italic = o.get('italic', italic)
            f.color.rgb = o.get('color', color)
    return tb


def box(slide, x, y, w, h, fill, radius=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    if radius:
        shp.adjustments[0] = 0.06
    return shp


def picture(slide, path, x, y, w, h):
    """Fit image into the box, keep aspect ratio, centre it."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def content_slide(kicker, title, source, notes=''):
    s = prs.slides.add_slide(BLANK)
    page = len(prs.slides)
    text(s, 0.6, 0.45, 12, 0.3, kicker.upper(), size=12, color=NEG, bold=True)
    text(s, 0.6, 0.78, 12.1, 0.95, title, size=30, color=NAVY, bold=True)
    text(s, 0.6, 7.0, 10.5, 0.3, source, size=10, color=MUTED)
    text(s, 11.9, 7.0, 0.85, 0.3, f'{page}', size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def stat(slide, x, y, w, big, label, color=NAVY, big_size=44):
    text(slide, x, y, w, 0.8, big, size=big_size, color=color, bold=True)
    text(slide, x, y + 0.85, w, 0.9, label, size=14, color=INK2)


def callout(slide, x, y, w, h, head, body, fill=TINT):
    box(slide, x, y, w, h, fill)
    text(slide, x + 0.25, y + 0.2, w - 0.5, h - 0.4,
         [[(head, {'bold': True, 'color': NAVY})], body], size=14, color=INK, spacing=6)


topics = K['topics']                       # sorted by share (desc)
T = {t['topic_name']: t for t in topics}
neg_share = K['n_negative'] / K['n_docs'] * 100
lukewarm_rec = K['recommend_by_rating']['3']
SRC = f"Source: clothing_reviews.xlsx · {K['n_docs']:,} reviews with text · analysis: customer_review_text_mining.ipynb"

# ---------------------------------------------------------------- 1 Title
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, NAVY, radius=False)
text(s, 0.8, 1.2, 11, 0.4, 'MANAGEMENT BRIEFING · TEXT MINING · OCTOBER 2026', size=14, color=RGBColor(0xF5, 0x9E, 0x7A), bold=True)
text(s, 0.8, 1.8, 11.5, 2.0, f"What {K['n_docs']:,} customer reviews tell us", size=48, color=WHITE, bold=True)
text(s, 0.8, 3.3, 10.5, 1.2,
     'A text mining analysis of women\'s apparel reviews: what customers talk about, '
     'what drives bad ratings, and where to act', size=20, color=RGBColor(0xD6, 0xDC, 0xE8))
text(s, 0.8, 5.6, 11, 0.4, 'Tim Vetter · Lisa-Marie Schwarz · Batuhan Aras · Max Tenz · Daniele Demerji',
     size=16, color=WHITE)
text(s, 0.8, 6.05, 11, 0.4, 'Text Mining / Text Analytics · Hochschule Ludwigshafen · Python & Jupyter Notebook',
     size=13, color=RGBColor(0xA9, 0xB3, 0xC6))
s.notes_slide.notes_text_frame.text = (
    'We analysed the free text of the reviews with text mining methods from the course, '
    'not just the star ratings. All numbers come from the Jupyter notebook customer_review_text_mining.ipynb.')

# ---------------------------------------------------------------- 2 Executive summary
top1, top2 = topics[0], topics[1]
s = content_slide('Executive summary',
                  'Customers are satisfied, but fit, cut and fabric drive the bad reviews',
                  SRC,
                  'Three numbers to remember: satisfaction is high; one in five reviews mentions a size problem; '
                  'and the topic model shows what unhappy customers complain about.')
cols = [
    (f"{K['recommend_share']:.0f}%", f"recommend the item · average {K['avg_rating']}★ · only {neg_share:.0f}% of reviews are negative (1–2★)", NAVY),
    (f"{K['size_mismatch_share']:.0f}%", f"of all reviews report a size mismatch (“runs large”, “size down”…), even {K['size_mismatch_share_pos']:.0f}% of the happy ones", NEG),
    (f"{top1['share'] + top2['share']:.0f}%", f"of critical (1–3★) reviews are about cut & shape ({top1['share']:.0f}%) or sizing ({top2['share']:.0f}%) – found by LDA topic modeling", NEG),
]
for i, (big, lab, c) in enumerate(cols):
    x = 0.6 + i * 4.1
    box(s, x, 2.0, 3.8, 2.2, TINT)
    stat(s, x + 0.3, 2.2, 3.3, big, lab, color=c, big_size=48)
text(s, 0.6, 4.75, 12, 0.4, 'Where to act', size=18, color=NAVY, bold=True)
acts = [('Fit guidance', 'per-item size advice from review text'),
        ('Cut & silhouette', 'feed "unflattering" reviews back to design'),
        ('Fabric & workmanship', 'opacity, weight and seam checks'),
        ('Product pages', 'photos and specs that match reality')]
for i, (h, b) in enumerate(acts):
    x = 0.6 + i * 3.08
    text(s, x, 5.25, 2.9, 1.2, [[(f'{i + 1} · {h}', {'bold': True, 'color': NAVY})], b], size=14, color=INK2, spacing=3)

# ---------------------------------------------------------------- 3 Approach
s = content_slide('Our approach',
                  'From 22,641 free-text reviews to numbers management can act on',
                  'Methods from the course: Text Mining · Information Retrieval · Natural Language Processing',
                  'This is the pipeline from the first lecture. Each step is a section in the notebook. '
                  'Key point: we did not define keyword lists in advance – the topics were discovered by the LDA algorithm.')
steps = [('1', 'Corpus', 'Title + review text\n22,641 documents\n1.36 M tokens'),
         ('2', 'Text metrics', 'Length, word length,\nvocabulary, Flesch\nreadability'),
         ('3', 'Preprocessing', 'Tokenization, stop-\nwords, lemmatization\n(NLTK)'),
         ('4', 'Frequencies', 'Words, n-grams,\nword clouds,\nco-occurrence'),
         ('5', 'TF-IDF', 'Document-term matrix,\nterms that separate\nbad from good'),
         ('6', 'Sentiment', 'VADER lexicon score\nvalidated against\nstar rating'),
         ('7', 'Topic model', 'LDA on 5,193\ncritical reviews,\n8 themes')]
w = 1.66
for i, (n, h, b) in enumerate(steps):
    x = 0.6 + i * (w + 0.12)
    fill = NAVY if i == 6 else TINT
    box(s, x, 2.1, w, 3.2, fill)
    c_head = WHITE if i == 6 else NAVY
    c_body = RGBColor(0xD6, 0xDC, 0xE8) if i == 6 else INK2
    text(s, x + 0.18, 2.3, w - 0.3, 0.6, n, size=28, color=NEG if i != 6 else RGBColor(0xF5, 0x9E, 0x7A), bold=True)
    text(s, x + 0.18, 3.0, w - 0.3, 0.5, h, size=16, color=c_head, bold=True)
    text(s, x + 0.18, 3.55, w - 0.25, 1.6, b, size=13, color=c_body)
text(s, 0.6, 5.65, 12.1, 1.0,
     [[('Reading the results: ', {'bold': True, 'color': NAVY}),
       ('negative = 1–2★, neutral = 3★, positive = 4–5★. Shares are “% of reviews that mention it”. '
        'Topic names are our interpretation of the words the algorithm grouped together.', {})]],
     size=14, color=INK2)

# ---------------------------------------------------------------- 4 Ratings
s = content_slide('Overall picture',
                  'Most customers give five stars – the 3★ group is the one to win',
                  SRC + ' · rating and recommendation columns',
                  'Before mining the text: the structured data. 3-star reviewers are split – only about 4 in 10 still recommend.')
picture(s, 'figures/01_rating_distribution.png', 0.6, 1.9, 7.4, 4.8)
stat(s, 8.6, 2.0, 4.2, f"{K['avg_rating']}★", 'average rating')
stat(s, 8.6, 3.55, 4.2, f"{K['recommend_share']:.0f}%", 'would recommend the item')
callout(s, 8.6, 5.05, 4.2, 1.45, f"Only {lukewarm_rec:.0f}% of 3★ reviewers recommend",
        f"{K['n_neutral']:,} lukewarm customers – their texts show what holds them back.")

# ---------------------------------------------------------------- 5 Text to numbers + word clouds
s = content_slide('Word frequencies',
                  'Happy and unhappy customers use the same words',
                  SRC + ' · after stop-word removal and lemmatization',
                  'Word clouds show what customers talk about: dress, fit, size, fabric, colour in both groups. '
                  'To see what separates good from bad, we need TF-IDF / log-odds (next slide).')
picture(s, 'figures/03_wordclouds.png', 0.5, 1.7, 12.3, 4.15)
text(s, 6.9, 5.95, 5.9, 0.95,
     [[('Text as numbers: ', {'bold': True, 'color': NAVY}),
       (f"{K['avg_words']:.0f} words per review on average, {K['avg_word_len']:.1f} letters per word, "
        f"Flesch reading ease {K['avg_flesch']:.0f} (“easy”), vocabulary of {K['vocabulary']:,} distinct words. "
        "3‑4★ reviews are the longest – customers explain trade-offs.", {})]], size=14, color=INK2)
text(s, 0.6, 5.95, 5.9, 0.95,
     [[('Frequency alone is not enough: ', {'bold': True, 'color': NAVY}),
       ('dress, fit, size, fabric and color dominate both sides. Positive adds love, perfect, flattering, soft; '
        'negative adds back (“sending it back”), material, quality.', {})]],
     size=14, color=INK2)

# ---------------------------------------------------------------- 6 Distinctive terms
nt = list(K['neg_terms'].items())
s = content_slide('TF-IDF & log-odds',
                  'Bad reviews are about cheap quality and unflattering shapes',
                  SRC + ' · terms in ≥ 50 negative / ≥ 300 positive reviews; pure emotion words removed',
                  'Each bar: how many times more often the term appears in a 1-2 star review than in a 4-5 star review. '
                  'Purely emotional words like "disappointed" were removed so the chart shows product drivers.')
picture(s, 'figures/06_distinctive_terms.png', 0.5, 1.85, 8.7, 4.95)
sh = K['neg_terms_share']
callout(s, 9.6, 2.3, 3.2, 1.75, 'Quality',
        f"“poor quality” {nt[0][1]:.0f}× and “cheap” "
        f"{K['neg_terms'].get('cheap', 0):.0f}× more frequent in bad reviews", fill=NEG_TINT)
callout(s, 9.6, 4.35, 3.2, 1.75, 'Shape',
        f"“unflattering” in {sh['unflattering']['neg %']:.0f}% of bad vs {sh['unflattering']['pos %']:.1f}% of good reviews; "
        '“shapeless”, “huge”, “tent”', fill=NEG_TINT)

# ---------------------------------------------------------------- 7 Sizing
sz = K['sizing']
s = content_slide('N-gram analysis: sizing',
                  'One in five reviews reports a size mismatch – items run large',
                  SRC + ' · bigram/trigram patterns such as “runs (a bit) large”, “size down”',
                  'N-grams turn sizing language into a measurable signal. "Runs large" and "size down" '
                  'clearly outnumber "runs small" and "size up". Even satisfied customers warn others about the size.')
picture(s, 'figures/04_sizing_ngrams.png', 0.5, 1.85, 7.8, 4.95)
stat(s, 8.9, 2.0, 3.9, f"{K['size_mismatch_share']:.0f}%", 'of reviews mention a size mismatch', color=NEG)
stat(s, 8.9, 3.6, 3.9, f"{K['size_mismatch_share_pos']:.0f}%", 'of 4–5★ reviews still warn about the size', color=NEG)
callout(s, 8.9, 5.1, 3.9, 1.6, f"“true to size” only {sz['true to size']:.0f}%",
        'Customers need size advice – and they are already writing it for us.')

# ---------------------------------------------------------------- 8 LDA topics
s = content_slide('Topic modeling (LDA)',
                  'The algorithm finds eight complaint themes – cut and sizing lead',
                  f"LDA (scikit-learn), k = 8, on {sum(t['reviews'] for t in topics):,} reviews with 1–3★ · dominant topic per review · words in brackets = top terms",
                  'LDA groups words that occur together into topics. We read the top words and example reviews '
                  'and named each topic. Cut & shape and sizing together account for almost half of critical reviews.')
picture(s, 'figures/08_lda_topics.png', 0.5, 1.8, 12.3, 5.05)

# ---------------------------------------------------------------- 9 Department
dt = K['dept_topic']
def dshare(prefix, topic):
    key = next(k for k in dt if k.startswith(prefix))
    return dt[key][topic]
s = content_slide('Topics by department',
                  'Beyond cut and sizing, each department has its own weak spot',
                  'Share of each department\'s 1–3★ reviews by dominant LDA topic (%) · Trend: small sample',
                  'Read across a row: Dresses – the cut; Bottoms – pant and jeans fit; Jackets and Intimates – workmanship; '
                  'Tops – washing and fabric. This tells each category team where to start.')
picture(s, 'figures/09_topics_by_department.png', 0.5, 1.8, 8.9, 5.0)
rows = [('Dresses', f"cut & shape {dshare('Dresses', 'Unflattering cut & shape'):.0f}%"),
        ('Bottoms', f"pants & jeans fit {dshare('Bottoms', 'Pants & jeans fit'):.0f}%"),
        ('Jackets', f"workmanship {dshare('Jackets', 'Workmanship: buttons, zips, seams'):.0f}%"),
        ('Intimate', f"workmanship {dshare('Intimate', 'Workmanship: buttons, zips, seams'):.0f}%, fabric {dshare('Intimate', 'Cheap fabric for the price'):.0f}%"),
        ('Tops', f"sweaters & wash {dshare('Tops', 'Sweaters: boxy, shrinks in the wash'):.0f}%, fabric {dshare('Tops', 'Cheap fabric for the price'):.0f}%")]
box(s, 9.7, 1.9, 3.1, 4.85, TINT)
text(s, 9.95, 2.1, 2.7, 0.4, 'Above-average topic', size=16, color=NAVY, bold=True)
for i, (d, v) in enumerate(rows):
    text(s, 9.95, 2.65 + i * 0.8, 2.7, 0.75, [[(d, {'bold': True, 'color': INK})], v], size=14, color=INK2)

# ---------------------------------------------------------------- 10 VADER
s = content_slide('Sentiment analysis (VADER)',
                  'Text sentiment follows the stars – but cannot replace them',
                  SRC + f" · VADER compound score; Spearman ρ with rating = {K['vader_corr']}",
                  'VADER is a dictionary-based sentiment model. It rises with the rating, but it scores 70% of 1-2 star reviews as positive, '
                  'because disappointed customers write "I wanted to love this, beautiful color, but...". '
                  'Lesson: the value of the text is in WHY customers are unhappy, not in a sentiment score.')
picture(s, 'figures/07_vader_vs_rating.png', 0.5, 1.85, 8.0, 4.9)
stat(s, 9.0, 2.0, 3.8, f"{K['vader_pos_in_neg']:.0f}%", 'of 1–2★ reviews are scored “positive” by VADER', color=NEG)
callout(s, 9.0, 3.75, 3.8, 2.35, '“I wanted to love this…”',
        'Unhappy customers praise colour and idea before the decisive “but it fits like a tent”. '
        'A general sentiment lexicon misses this – use topics and terms, not a sentiment score, for reporting.')

# ---------------------------------------------------------------- 11 Age
age = K['age']
s = content_slide('By age group',
                  'Satisfaction rises with age – the complaints are the same for everyone',
                  SRC + ' · age as stated by the reviewer',
                  'All age groups name the same LDA topics in roughly the same proportions. '
                  'So the fixes are product fixes, not age-specific campaigns.')
tbl_rows = len(age) + 1
tb = s.shapes.add_table(tbl_rows, 5, Inches(0.6), Inches(1.95), Inches(7.6), Inches(0.5 * tbl_rows)).table
heads = ['Age', 'Share of reviewers', 'Avg ★', 'Negative', 'Recommend']
for j, h in enumerate(heads):
    c = tb.cell(0, j); c.text = h
for i, a in enumerate(age, start=1):
    vals = [a['age_group'], f"{a['share']:.0f}%", f"{a['avg_rating']:.2f}", f"{a['negative']:.1f}%", f"{a['recommend']:.0f}%"]
    for j, v in enumerate(vals):
        tb.cell(i, j).text = v
for i in range(tbl_rows):
    for j in range(5):
        c = tb.cell(i, j)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (TINT if i % 2 == 0 else WHITE)
        for p in c.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            for r in p.runs:
                r.font.name = FONT; r.font.size = Pt(14)
                r.font.color.rgb = WHITE if i == 0 else INK
                r.font.bold = i == 0
at = K['age_topic']
cut_young, cut_old = at['18-29']['Unflattering cut & shape'], at['60+']['Unflattering cut & shape']
size_rng = [at[g]['Sizing: runs small / large'] for g in at]
stat(s, 8.9, 2.0, 3.9, f"{age[-1]['recommend']:.0f}%", 'of reviewers aged 60+ recommend – the highest group')
callout(s, 8.9, 3.75, 3.9, 2.0, 'Same problems at every age',
        f"Sizing is {min(size_rng):.0f}–{max(size_rng):.0f}% of critical reviews in every age group. "
        f"Cut & shape matters slightly more to younger customers ({cut_young:.0f}% at 18–29 vs {cut_old:.0f}% at 60+).")

# ---------------------------------------------------------------- 12 Recommendations
s = content_slide('Recommendations',
                  'Four actions – each with a KPI we can track from the review text',
                  'KPIs are measured with the same notebook – re-run it on each new review export',
                  'Each recommendation links to a finding and a measurable KPI. Because the analysis is a notebook, '
                  'these KPIs can be refreshed every quarter at no extra cost.')
recs = [
    ('1', 'Fix fit guidance', 'Per-item fit notes (“runs large – size down”) and a size advisor fed by review n-grams. Start with Dresses and Tops.',
     f"KPI: size-mismatch reviews, today {K['size_mismatch_share']:.0f}%"),
    ('2', 'Review cuts & silhouettes', 'Share the “unflattering / shapeless / boxy” reviews with design every season; rework repeat offenders.',
     f"KPI: cut & shape topic, today {T['Unflattering cut & shape']['share']:.0f}% of critical reviews"),
    ('3', 'Tighten fabric & workmanship checks', 'Opacity, fabric weight, wash tests and seam/zip checks – especially Tops, Intimates, Jackets.',
     f"KPI: fabric + workmanship + washing topics, today {T['Cheap fabric for the price']['share'] + T['Workmanship: buttons, zips, seams']['share'] + T['Sweaters: boxy, shrinks in the wash']['share']:.0f}%"),
    ('4', 'Make product pages match reality', 'Show fabric, opacity, length and the model’s size; re-shoot items with “not like the picture” reviews.',
     f"KPI: see-through / photo topic, today {T['See-through / differs from photo']['share']:.0f}%"),
]
for i, (n, h, b, kpi) in enumerate(recs):
    x = 0.6 + (i % 2) * 6.15
    y = 1.95 + (i // 2) * 2.45
    box(s, x, y, 5.95, 2.25, TINT)
    text(s, x + 0.25, y + 0.2, 0.6, 0.7, n, size=32, color=NEG, bold=True)
    text(s, x + 0.9, y + 0.25, 4.85, 0.45, h, size=18, color=NAVY, bold=True)
    text(s, x + 0.9, y + 0.75, 4.85, 0.9, b, size=14, color=INK2)
    text(s, x + 0.9, y + 1.68, 4.85, 0.4, kpi, size=13, color=NEG, bold=True)

# ---------------------------------------------------------------- 13 Appendix
s = content_slide('Appendix · method and limits',
                  'How the analysis was done and what to keep in mind',
                  'All code, outputs and charts: customer_review_text_mining.ipynb · github.com/Tim-vt22/MBA-Text-Mining',
                  'For questions on method. The notebook is fully reproducible.')
text(s, 0.6, 1.95, 5.9, 4.8, [
    [('Data & tools', {'bold': True, 'color': NAVY, 'size': 18})],
    f"clothing_reviews.xlsx: {K['rows_file']:,} rows → {K['n_rated']:,} rated → {K['n_docs']:,} with text (title + review = one document)",
    'Python in Jupyter: pandas, NLTK (tokenization, stop-words, WordNet lemmatizer), scikit-learn (CountVectorizer, TF-IDF, LDA), vaderSentiment, textstat, wordcloud',
    'Negations kept for n-grams (“not flattering”, “too big”)',
    'LDA: 1,111-term vocabulary, k = 8 chosen for interpretability, fixed random seed',
], size=14, color=INK2, spacing=8)
text(s, 7.0, 1.95, 5.8, 4.8, [
    [('Limits', {'bold': True, 'color': NAVY, 'size': 18})],
    'No review dates in the export – no trends over time yet',
    'Review texts are cut at ~500 characters',
    'One dominant topic per review for reporting; real reviews mix topics. Topic names are our interpretation',
    'VADER is a general lexicon, not tuned to fashion language',
    'Trend department (118 reviews with text) too small for firm conclusions',
], size=14, color=INK2, spacing=8)

prs.save(OUT)
print('saved', OUT, 'slides:', len(prs.slides))
